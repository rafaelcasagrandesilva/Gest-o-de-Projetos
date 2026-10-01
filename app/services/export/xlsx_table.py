"""Padrão ÚNICO de Excel do SGC: cada aba é uma Tabela do Excel no estilo azul-marinho.

Aprovado pelo usuário (out/2026) para TODOS os relatórios:

- cabeçalho azul-marinho com texto branco em negrito, centralizado;
- linhas alternadas bem claras, borda fina cinza-claro;
- valores NUMÉRICOS em formato Contábil R$ e datas reais — nunca o texto "R$ 1.234,56", que o
  Excel marca com o triângulo verde e não soma;
- linha de Total da Tabela com SUBTOTAL (acompanha o filtro), em negrito com borda dupla;
- linha 1 e coluna A congeladas, sem linhas de grade, SEM quebra de texto (altura uniforme).

Os builders de `builders.py` são camadas finas sobre `write_table_sheet`; um relatório novo
herda o padrão sem fazer nada.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table as XlTable
from openpyxl.worksheet.table import TableColumn, TableStyleInfo

# Formato "Contábil" do Excel em R$: zero aparece como "-".
ACCOUNTING_BRL_FMT = '_-"R$"\\ * #,##0.00_-;\\-"R$"\\ * #,##0.00_-;_-"R$"\\ * "-"??_-;_-@_-'
PERCENT_FMT = "0.00%"
DATE_FMT = "DD/MM/YYYY"

_NAVY = "1F3864"
_HEADER_FILL = PatternFill(start_color=_NAVY, end_color=_NAVY, fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_GRID_SIDE = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_GRID_SIDE, right=_GRID_SIDE, top=_GRID_SIDE, bottom=_GRID_SIDE)
_TOTAL_BORDER = Border(
    left=_GRID_SIDE, right=_GRID_SIDE, top=Side(style="double", color="1F1F1F"), bottom=_GRID_SIDE
)
_TABLE_STYLE = "TableStyleLight1"  # listras cinza bem claras
_RED = "C00000"

# Funções da linha de total → código do SUBTOTAL (o 1xx ignora linhas escondidas pelo filtro).
_SUBTOTAL_CODES = {"sum": 109, "count": 103}

_BRL_TEXT = re.compile(r"^(-)?\s*R\$\s*(-)?\s*(\d{1,3}(?:\.\d{3})*|\d+),(\d{2})$")
_PCT_TEXT = re.compile(r"^(-?\d+(?:,\d+)?)\s*%$")
_DATE_TEXT = re.compile(r"^\d{2}/\d{2}/\d{4}$")


@dataclass
class SheetSpec:
    """Uma aba. Colunas 1-based.

    `totals`: {coluna: "sum" | "count" | valor fixo}. Vazio = Tabela sem linha de total.
    `parse_text`: converte "R$ 1.234,56" → número, "12,5%" → percentual e "dd/mm/aaaa" → data,
    para relatórios antigos que ainda montam as células como texto formatado.
    """

    title: str
    headers: Sequence[str]
    rows: Sequence[Sequence[Any]]
    money_columns: frozenset[int] = frozenset()
    date_columns: frozenset[int] = frozenset()
    percent_columns: frozenset[int] = frozenset()
    totals: dict[int, Any] = field(default_factory=dict)
    total_label: str = "Total"
    bold_columns: frozenset[int] = frozenset()
    negative_red: bool = False
    parse_text: bool = False
    # Soma automática em toda coluna de valor (inclusive as descobertas pelo `parse_text`);
    # `totals` sobrepõe coluna a coluna.
    auto_totals: bool = False
    freeze: str = "B2"


def build_workbook_bytes(sheets: Sequence[SheetSpec]) -> bytes:
    wb = Workbook()
    for i, spec in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = _sheet_title(spec.title, i, {s.title for s in wb.worksheets if s is not ws})
        write_table_sheet(ws, spec, table_name=f"Tabela{i + 1}")
    # Fórmulas de total saem sem valor em cache: o Excel recalcula ao abrir.
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def write_table_sheet(ws, spec: SheetSpec, *, table_name: str) -> None:
    headers = _unique_headers(spec.headers)
    ncols = len(headers)
    if ncols == 0:
        return
    rows = [list(r) + [None] * (ncols - len(r)) for r in spec.rows] or [[None] * ncols]
    money = set(spec.money_columns)
    dates = set(spec.date_columns)
    percents = set(spec.percent_columns)
    cell_formats: dict[tuple[int, int], str] = {}

    if spec.parse_text:
        rows, money, dates, percents, cell_formats = _parse_text_cells(rows, money, dates, percents)

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    ws.row_dimensions[1].height = 22

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row[:ncols], start=1):
            if c_idx in dates:
                val = _as_date(val)
            if val == "":
                val = None
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(vertical="top")
            cell.border = _BORDER
            fmt = _column_format(c_idx, money, dates, percents) or cell_formats.get((r_idx, c_idx))
            if fmt and val is not None:
                cell.number_format = fmt
            _font(cell, bold=c_idx in spec.bold_columns, red=spec.negative_red)

    totals = dict(spec.totals)
    if spec.auto_totals:
        totals = {**{c: "sum" for c in money if c != 1}, **totals}
    last_data = len(rows) + 1
    has_totals = bool(totals)
    total_row = last_data + 1 if has_totals else last_data
    columns: list[TableColumn] = []
    for c, name in enumerate(headers, start=1):
        col = TableColumn(id=c, name=name)
        if has_totals:
            cell = ws.cell(row=total_row, column=c)
            cell.border = _TOTAL_BORDER
            func = totals.get(c)
            if func in _SUBTOTAL_CODES:
                col.totalsRowFunction = func
                cell.value = f"=SUBTOTAL({_SUBTOTAL_CODES[func]},{table_name}[{_table_ref(name)}])"
            elif func is not None:
                cell.value = func  # valor fixo (ex.: margem %, que não é soma)
            elif c == 1:
                col.totalsRowLabel = spec.total_label
                cell.value = spec.total_label
            fmt = _column_format(c, money, dates, percents)
            if fmt and func is not None and func != "count":
                cell.number_format = fmt
            _font(cell, bold=True, red=spec.negative_red)
        columns.append(col)

    table = XlTable(displayName=table_name, ref=f"A1:{get_column_letter(ncols)}{total_row}")
    table.tableColumns = columns
    table.tableStyleInfo = TableStyleInfo(name=_TABLE_STYLE, showRowStripes=True)
    # Botões de filtro no cabeçalho; a linha de total fica FORA do autofiltro (senão o filtro
    # a esconderia).
    table.autoFilter = AutoFilter(ref=f"A1:{get_column_letter(ncols)}{last_data}")
    if has_totals:
        table.totalsRowCount = 1
    ws.add_table(table)

    for c in range(1, ncols + 1):
        longest = len(headers[c - 1]) + 4  # espaço da seta do filtro
        for r_idx in range(2, last_data + 1):
            v = ws.cell(row=r_idx, column=c).value
            if v is None:
                continue
            if isinstance(v, (int, float)) and c in money:
                text = f"R$ {v:,.2f}"
            elif isinstance(v, (date, datetime)):
                text = "00/00/0000"
            else:
                text = str(v)
            longest = max(longest, min(len(text), 50))
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 10), 52)

    ws.freeze_panes = spec.freeze
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------------------------------------


def _font(cell, *, bold: bool, red: bool) -> None:
    is_negative = red and isinstance(cell.value, (int, float)) and cell.value < 0
    if bold or is_negative:
        cell.font = Font(bold=bold, color=_RED if is_negative else None)


def _column_format(c: int, money: set[int], dates: set[int], percents: set[int]) -> str | None:
    if c in money:
        return ACCOUNTING_BRL_FMT
    if c in percents:
        return PERCENT_FMT
    if c in dates:
        return DATE_FMT
    return None


def _parse_text_cells(rows, money, dates, percents):
    """Converte células de texto formatado em valores de verdade.

    Uma coluna vira de valor/data/percentual quando TODAS as células preenchidas convertem
    (traço "—" conta como vazio). Células soltas que convertem — como a coluna "Valor" de uma
    aba de indicadores, que mistura R$ e % — ganham o formato só nelas.
    """
    ncols = max((len(r) for r in rows), default=0)
    out = [list(r) for r in rows]
    kinds: dict[int, set[str]] = {c: set() for c in range(1, ncols + 1)}
    cell_formats: dict[tuple[int, int], str] = {}
    for r_idx, row in enumerate(out, start=2):
        for c_idx, val in enumerate(row, start=1):
            if not isinstance(val, str):
                if val is not None:
                    kinds[c_idx].add("number" if isinstance(val, (int, float)) else "other")
                continue
            text = val.strip()
            if text in ("", "—", "-"):
                row[c_idx - 1] = None
                continue
            kind, parsed = _parse_one(text)
            kinds[c_idx].add(kind)
            if kind != "text":
                row[c_idx - 1] = parsed
                cell_formats[(r_idx, c_idx)] = {
                    "money": ACCOUNTING_BRL_FMT, "percent": PERCENT_FMT, "date": DATE_FMT
                }[kind]
    for c, found in kinds.items():
        if found == {"money"} or found == {"money", "number"}:
            money.add(c)
        elif found == {"date"}:
            dates.add(c)
        elif found == {"percent"}:
            percents.add(c)
    return out, money, dates, percents, cell_formats


def _parse_one(text: str) -> tuple[str, Any]:
    m = _BRL_TEXT.match(text)
    if m:
        neg = bool(m.group(1) or m.group(2))
        value = float(f"{m.group(3).replace('.', '')}.{m.group(4)}")
        return "money", -value if neg else value
    m = _PCT_TEXT.match(text)
    if m:
        return "percent", round(float(m.group(1).replace(",", ".")) / 100, 6)
    if _DATE_TEXT.match(text):
        try:
            return "date", datetime.strptime(text, "%d/%m/%Y").date()
        except ValueError:
            return "text", text
    return "text", text


def _unique_headers(headers: Sequence[Any]) -> list[str]:
    """Nomes de coluna de Tabela do Excel precisam ser únicos e não vazios."""
    seen: dict[str, int] = {}
    out: list[str] = []
    for i, h in enumerate(headers, start=1):
        name = str(h).strip().replace("\n", " ") or f"Coluna{i}"
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name} {seen[key]}"
        else:
            seen[key] = 1
        out.append(name)
    return out


def _table_ref(name: str) -> str:
    """Escapa o nome da coluna para referência estruturada (Tabela1[Coluna])."""
    return "".join(f"'{ch}" if ch in "[]#'" else ch for ch in name)


def _as_date(val: Any) -> Any:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date) or val in (None, ""):
        return val if val != "" else None
    text = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return val


def _sheet_title(title: str, index: int, taken: set[str]) -> str:
    base = (re.sub(r"[\[\]:*?/\\]", "-", str(title or "")).strip() or f"Aba {index + 1}")[:31]
    name, n = base, 2
    while name in taken:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    return name
