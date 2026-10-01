"""Relatórios em formato Tabela do Excel (Folha e Contas a Pagar)."""

from __future__ import annotations

import io
import unittest
from datetime import date

from openpyxl import load_workbook


class TableXlsxBuilderTests(unittest.TestCase):
    def _wb(self):
        from app.services.export.builders import build_table_xlsx_bytes

        raw = build_table_xlsx_bytes(
            headers=["Nome", "Vencimento", "Valor", "Total"],
            rows=[["A", "2026-10-10", 1234.5, 1234.5], ["B", None, None, 10.0]],
            money_columns=frozenset({3, 4}),
            date_columns=frozenset({2}),
            totals={3: "sum", 4: "sum"},
            bold_columns=frozenset({4}),
        )
        return load_workbook(io.BytesIO(raw))

    def test_table_with_filter_aware_totals(self):
        ws = self._wb().active
        table = ws.tables["Tabela1"]
        self.assertEqual(table.ref, "A1:D4")
        self.assertEqual(table.totalsRowCount, 1)
        self.assertEqual(table.autoFilter.ref, "A1:D3")  # total fora do filtro
        self.assertEqual(ws["A4"].value, "Total")
        self.assertEqual(ws["C4"].value, "=SUBTOTAL(109,Tabela1[Valor])")

    def test_numbers_dates_and_layout(self):
        ws = self._wb().active
        self.assertEqual(ws["C2"].value, 1234.5)  # número, não texto "R$ 1.234,50"
        self.assertIn('"R$"', ws["C2"].number_format)
        self.assertEqual(ws["B2"].value.date(), date(2026, 10, 10))
        self.assertEqual(ws.freeze_panes, "B2")
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertFalse(any(c.alignment.wrap_text for row in ws.iter_rows() for c in row))
        self.assertEqual(ws["A1"].alignment.horizontal, "center")
        self.assertTrue(ws["D2"].font.b)

    def test_cap_report_uses_table(self):
        from app.services.operational_report_export import render_operational_report_bytes

        payload = {"rows": [{"nome": "X", "vencimento": "2026-10-10", "valor_final": 100.0, "saldo": 100.0}]}
        raw, _, _ = render_operational_report_bytes("payables_detailed", payload, "xlsx", None)
        ws = load_workbook(io.BytesIO(raw)).active
        self.assertIn("Tabela1", ws.tables)
        headers = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=headers.index("Valor final") + 1).value, 100.0)


class StandardEverywhereTests(unittest.TestCase):
    """Todo builder de Excel sai no padrão único (azul-marinho, Tabela, sem texto-moeda)."""

    def _check(self, raw: bytes):
        wb = load_workbook(io.BytesIO(raw))
        for ws in wb:
            self.assertTrue(ws.tables, ws.title)
            self.assertTrue(ws["A1"].fill.fgColor.rgb.endswith("1F3864"), ws.title)
            self.assertEqual(ws.freeze_panes, "B2")
            self.assertFalse(ws.sheet_view.showGridLines)
        return wb

    def test_legacy_text_builder_converts_money_dates_and_totals(self):
        from app.services.export.builders import build_xlsx_bytes

        raw = build_xlsx_bytes(
            headers=["Projeto", "Valor", "Vencimento", "Margem"],
            rows=[["A", "R$ 1.234,56", "10/10/2026", "12,50%"], ["B", "-R$ 34,56", "", "—"]],
            totals_row=["Totais", "R$ 1.200,00", "", ""],
        )
        ws = self._check(raw).active
        self.assertEqual(ws["B2"].value, 1234.56)
        self.assertEqual(ws["B3"].value, -34.56)
        self.assertEqual(ws["C2"].value.date(), date(2026, 10, 10))
        self.assertAlmostEqual(ws["D2"].value, 0.125)
        self.assertEqual(ws["D2"].number_format, "0.00%")
        # O total informado é a soma da coluna → vira SUBTOTAL (segue o filtro).
        self.assertEqual(ws["B4"].value, "=SUBTOTAL(109,Tabela1[Valor])")

    def test_legacy_fixed_total_is_kept_when_not_a_sum(self):
        from app.services.export.builders import build_xlsx_bytes

        raw = build_xlsx_bytes(
            headers=["Item", "Valor"], rows=[["A", "R$ 10,00"]], totals_row=["Total", "R$ 99,00"]
        )
        self.assertEqual(load_workbook(io.BytesIO(raw)).active["B3"].value, 99.0)

    def test_multisheet_and_projects_summary(self):
        from app.services.export.builders import (
            build_multisheet_operational_xlsx_bytes,
            build_projects_summary_xlsx_bytes,
        )

        wb = self._check(build_multisheet_operational_xlsx_bytes([
            {"title": "Lista", "headers": ["Nome", "Valor"], "rows": [["x", 1.0]],
             "money_columns": frozenset({2})},
            {"title": "Indicadores", "headers": ["Nome", "Valor"], "rows": [["y", 2.0]],
             "money_columns": frozenset({2}), "totals": False},
            {"title": "Vazia", "headers": ["Nome"], "rows": []},
        ]))
        self.assertEqual(wb["Lista"].tables["Tabela1"].totalsRowCount, 1)
        self.assertFalse(wb["Indicadores"].tables["Tabela2"].totalsRowCount)

        ws = self._check(build_projects_summary_xlsx_bytes(
            headers=["Projeto", "Receita", "Margem"],
            rows=[["A", 100.0, 0.1], ["B", -50.0, -0.2]],
            totals_row=["Total", 50.0, -0.05],
        )).active
        self.assertEqual(ws["C2"].number_format, "0.00%")
        self.assertEqual(ws["B3"].font.color.rgb[-6:], "C00000")  # negativo em vermelho
        self.assertEqual(ws["C4"].value, -0.05)  # margem total é fixa, não soma


if __name__ == "__main__":
    unittest.main()
